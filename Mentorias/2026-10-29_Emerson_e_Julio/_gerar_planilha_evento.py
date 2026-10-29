# -*- coding: utf-8 -*-
# ==============================================================================
#   PLANILHA — só os dias com comunicado (mentoria de 29/10/2026)
#   Lê CVM/dados/dias_com_evento.json → 02_Dias_Com_Evento.xlsx
# ==============================================================================
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
SAIDA = AQUI / "02_Dias_Com_Evento.xlsx"
FONTE, AZUL, AZUL_IN = "Arial", "1F3864", "0000FF"
CINZA, AMARELO, VERDE, VERMELHO = "F2F2F2", "FFF2CC", "E2EFDA", "FCE4E4"
BORDA = Border(bottom=Side(style="thin", color="BFBFBF"))

D = json.loads((RAIZ / "CVM" / "dados" / "dias_com_evento.json").read_text(encoding="utf-8"))
P, AG = D["papeis"], D["agrupado"]
TICK = list(P.keys())


def cab(ws, tit, lg, linha=1):
    for c, (r_, w) in enumerate(zip(tit, lg), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
        cel = ws.cell(row=linha, column=c, value=r_)
        cel.font = Font(name=FONTE, bold=True, size=10, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=AZUL)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = BORDA
    ws.row_dimensions[linha].height = 34
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def txt(ws, r, c, v, bold=False, size=10, cor="000000", wrap=False, align=None, fundo=None):
    cel = ws.cell(row=r, column=c, value=v)
    cel.font = Font(name=FONTE, bold=bold, size=size, color=cor)
    cel.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="center")
    if fundo:
        cel.fill = PatternFill("solid", fgColor=fundo)
    return cel


def val(ws, r, c, v, fmt=None, cor=AZUL_IN, fundo=None):
    cel = ws.cell(row=r, column=c, value=v)
    cel.font = Font(name=FONTE, size=10, color=cor)
    if fmt:
        cel.number_format = fmt
    if fundo:
        cel.fill = PatternFill("solid", fgColor=fundo)
    cel.alignment = Alignment(horizontal="right")
    return cel


def bloco_texto(ws, r, itens, ncols):
    for titulo, corpo in itens:
        txt(ws, r, 1, titulo, bold=True, size=10, cor=AZUL)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols); r += 1
        txt(ws, r, 1, corpo, wrap=True, size=9, cor="595959")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.row_dimensions[r].height = 44; r += 2
    return r


wb = Workbook()

# ═══ ABA 1 — A RESPOSTA ═════════════════════════════════════════════════════
ws = wb.active
ws.title = "A resposta"
txt(ws, 1, 1, "Olhar só os dias com comunicado melhora os índices?", bold=True, size=15, cor=AZUL)
txt(ws, 2, 1, "A hipótese era que sim. O resultado foi o contrário — e o motivo é o que interessa.",
    size=10, cor="595959")
cab(ws, ["Arranjo", "Todos os dias\n(ganho sobre o palpite fixo)", "p",
         "SÓ dias com comunicado\n(ganho sobre o palpite fixo)", "p", "Melhorou?"],
    [46, 19, 9, 22, 9, 12], linha=4)
r, r0 = 5, 5
for b, v in AG.items():
    a, e = v["todos_os_dias"], v["so_dias_com_evento"]
    txt(ws, r, 1, b)
    val(ws, r, 2, a["ganho_pp"], "+0.00;-0.00", fundo=VERMELHO if a["ganho_pp"] < 0 else VERDE)
    val(ws, r, 3, a["p_vs_majoritaria"], "0.000", cor="000000")
    val(ws, r, 4, e["ganho_pp"], "+0.00;-0.00", fundo=VERMELHO if e["ganho_pp"] < 0 else VERDE)
    val(ws, r, 5, e["p_vs_majoritaria"], "0.000", cor="000000")
    txt(ws, r, 6, "sim" if e["ganho_pp"] > a["ganho_pp"] else "NÃO", bold=True, align="center",
        cor="C00000" if e["ganho_pp"] <= a["ganho_pp"] else "006100")
    for c in range(1, 7):
        ws.cell(row=r, column=c).border = BORDA
    if (r - r0) % 2 == 1:
        for c in range(1, 7):
            if ws.cell(row=r, column=c).fill.fgColor.rgb in (None, "00000000"):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=CINZA)
    r += 1
n_ev = AG[list(AG)[0]]["so_dias_com_evento"]["n"]
ic = AG[list(AG)[0]]["so_dias_com_evento"]["ic95_meia_largura_pp"]
r += 1
r = bloco_texto(ws, r, [
    ("Como ler esta tabela",
     f"Os 10 papéis foram agrupados num só conjunto, o que dá {AG[list(AG)[0]]['todos_os_dias']['n']} "
     f"pregões no total e {n_ev} com comunicado — margem de erro de ±{ic} pontos, pequena o "
     f"bastante para medir. Negativo significa PIOR que o palpite fixo (quem só repete o lado mais "
     f"comum). Cada coluna é comparada com o palpite fixo DO SEU PRÓPRIO subconjunto — comparar com "
     f"o do outro seria erro."),
    ("O resultado",
     "Em todos os seis arranjos, olhar só os dias com comunicado PIOROU. E nos dois arranjos que "
     "usam apenas preço, o déficit chega a ser estatisticamente significativo (p=0,022 e p=0,003) — "
     "ou seja, nesses dias o modelo é confiavelmente pior que não fazer nada."),
    ("Um padrão que vale notar, mas não afirmar",
     "Compare as duas primeiras linhas (só preço: −4,48 e −5,82) com as quatro seguintes (com "
     "sentimento: −1,34 a −3,43). Nos dias com comunicado, o sentimento parece REDUZIR o erro, "
     "ainda que não produza acerto. É um padrão consistente, mas não testei a diferença entre os "
     "arranjos diretamente — fica como conjectura, não como achado."),
], 6)

# ═══ ABA 2 — POR PAPEL ══════════════════════════════════════════════════════
ws1 = wb.create_sheet("Por papel")
txt(ws1, 1, 1, "Por que os papéis isolados não permitem concluir nada", bold=True, size=13, cor=AZUL)
cab(ws1, ["Papel", "Pregões de teste", "Com comunicado", "%", "Margem de erro\n(±pontos)",
          "Ganho médio\ntodos os dias", "Ganho médio\nsó com comunicado", "Braços\nsignificativos"],
    [9, 14, 14, 8, 15, 15, 18, 13], linha=3)
r, r0 = 4, 4
for tk in TICK:
    p = P[tk]; res = p.get("resumo", {})
    txt(ws1, r, 1, tk, bold=True)
    val(ws1, r, 2, p["n_teste"], "#,##0", cor="000000")
    val(ws1, r, 3, p["n_com_evento"], "#,##0", cor="000000")
    val(ws1, r, 4, p["pct_com_evento"] / 100, "0.0%", cor="000000")
    if res.get("ic95_com_evento_pp") is None:
        txt(ws1, r, 5, "não avaliável", size=9, align="center", cor="C00000")
        for c in range(6, 9):
            txt(ws1, r, c, "—", align="center")
    else:
        val(ws1, r, 5, res["ic95_com_evento_pp"], "0.0", fundo=VERMELHO)
        val(ws1, r, 6, res["ganho_medio_todos_pp"], "+0.00;-0.00")
        val(ws1, r, 7, res["ganho_medio_com_evento_pp"], "+0.00;-0.00",
            fundo=VERMELHO if res["ganho_medio_com_evento_pp"] < res["ganho_medio_todos_pp"] else VERDE)
        txt(ws1, r, 8, f"{res['bracos_significativos_com_evento']} de 6", align="center", bold=True)
    for c in range(1, 9):
        ws1.cell(row=r, column=c).border = BORDA
    r += 1
r += 1
r = bloco_texto(ws1, r, [
    ("A coluna que explica tudo é a margem de erro",
     "Com 33 a 165 pregões de comunicado, a margem de erro fica entre ±7,6 e ±17,1 pontos "
     "percentuais. O efeito que procuramos é da ordem de 2 ou 3 pontos. É como tentar pesar uma "
     "carta numa balança de caminhão: o instrumento não tem precisão para a pergunta."),
    ("A ABEV3 saiu da análise",
     "Só 5 pregões com comunicado em todo o período de teste (0,9%). Abaixo do mínimo de 30 que "
     "estabeleci antes de rodar. Não foi excluída por dar resultado ruim — foi excluída por não "
     "haver o que medir."),
    ("O modelo especialista também não resolveu",
     "Testei treinar o modelo APENAS com dias de comunicado, para afastar a objeção de que o modelo "
     "geral tivesse aprendido num mundo em que 60% dos dias não têm comunicado. Piorou em 6 dos 7 "
     "papéis em que havia treino suficiente. Em BBDC4 e BPAC11 nem foi possível treinar — menos de "
     "100 dias de comunicado no treino."),
], 8)

wb.save(SAIDA)
print(f"gravado: {SAIDA}  ({len(wb.sheetnames)} abas)")
