# -*- coding: utf-8 -*-
# ==============================================================================
#   PLANILHA — comparativos de projeção de direção (itens 2 a 7 do pedido do
#   Prof. Julio, 30/09/2026). Lê CVM/dados/preco_e_cvm_sem_noticia.json
#   (itens 2, 4, 5, 6 — executados em 01/10/2026 nesta máquina) e
#   CVM/dados/pesquisa_com_cvm.json (item 3 — execução anterior, outra janela
#   de teste; reaproveitado, não recalculado). Grava
#   17_Comparativos_Direcao.xlsx, nesta mesma pasta.
#
#   Os itens 3 e 7 do pedido (que precisam do corpus de notícias, ausente
#   nesta máquina) aparecem marcados como tal — item 3 com o número já
#   existente de outra execução; item 7 como bloqueado, sem número algum.
#
#   Quatro abas: Resumo por item | Direção - métricas completas |
#   McNemar - pares | Notas e limitações.
#
#   Valores de acurácia, AUC e p-valor são INPUT (resultado medido por um
#   modelo treinado) — cor azul, com a fonte (arquivo + chave) ao lado.
#   "Ganho em pontos" e "Passa no teste?" são FÓRMULA, nunca digitados.
# ==============================================================================
import json
from pathlib import Path

from scipy.stats import binomtest
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
SAIDA = AQUI / "17_Comparativos_Direcao.xlsx"

FONTE_NOME = "Arial"
AZUL_CAB = "1F3864"
AZUL_INPUT = "0000FF"
CINZA_CLARO = "F2F2F2"
AMARELO = "FFF2CC"
BORDA = Border(bottom=Side(style="thin", color="BFBFBF"))

d20 = json.loads((RAIZ / "CVM" / "dados" / "preco_e_cvm_sem_noticia.json").read_text(encoding="utf-8"))
d18 = json.loads((RAIZ / "CVM" / "dados" / "pesquisa_com_cvm.json").read_text(encoding="utf-8"))

FONTE_20 = "CVM/dados/preco_e_cvm_sem_noticia.json (executado em 01/10/2026, PETR4, teste=579 pregões)"
FONTE_18 = "CVM/dados/pesquisa_com_cvm.json (execução anterior, outra janela, PETR4, teste=653 pregões)"
FONTE_20_CURTA = "Hoje (script 20)"
FONTE_18_CURTA = "Anterior (script 18)"

# pesquisa_com_cvm.json só trazia p contra 50%; recalcula contra a classe
# majoritária (0,5314), que é o adversário correto — ver conversa de 21/09/2026
_item3 = d18["direcao"]["Data Fusion: preços + notícia  [O QUE TÍNHAMOS ANTES]"]
_n3, _maj3 = d18["n"]["teste"], d18["majoritaria_teste"]
for _mn, _m in _item3.items():
    _k = round(_m["acuracia"] * _n3)
    _m["p_binomial_vs_majoritaria"] = binomtest(_k, _n3, _maj3).pvalue


def cab(ws, titulos, larguras, linha=1):
    for c, (rot, larg) in enumerate(zip(titulos, larguras), start=1):
        col = get_column_letter(c)
        ws.column_dimensions[col].width = larg
        cel = ws.cell(row=linha, column=c, value=rot)
        cel.font = Font(name=FONTE_NOME, bold=True, size=10, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=AZUL_CAB)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = BORDA
    ws.row_dimensions[linha].height = 32
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def linha_zebra(ws, r, linha0):
    if (r - linha0) % 2 == 1:
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=CINZA_CLARO)


def input_cel(ws, r, c, valor, fmt=None, azul=True):
    cel = ws.cell(row=r, column=c, value=valor)
    cel.font = Font(name=FONTE_NOME, size=10, color=AZUL_INPUT if azul else "000000")
    if fmt:
        cel.number_format = fmt
    cel.alignment = Alignment(horizontal="right" if fmt else "left")
    return cel


def texto_cel(ws, r, c, valor, bold=False):
    cel = ws.cell(row=r, column=c, value=valor)
    cel.font = Font(name=FONTE_NOME, size=10, bold=bold)
    return cel


# ──────────────────────────────────────────────────────────────────────────────
wb = Workbook()

# ═══ ABA 1 — RESUMO POR ITEM (1 a 7) ═══════════════════════════════════════
ws1 = wb.active
ws1.title = "Resumo por item (1-7)"
titulos = ["Item", "Pedido do Prof. Julio", "Status", "Melhor acerto\n(SVM / XGB)",
           "Classe\nmajoritária", "p vs palpite fixo\n(melhor dos dois)", "Passa\n(p<0,05)?"]
larg = [7, 38, 30, 16, 13, 18, 11]
cab(ws1, titulos, larg)

ITENS = [
    (1, "Média de ações por investidor; investidores por ação",
     "Executado — ver 16_Distribuicao_de_Capital.xlsx", None, None, None),
    (2, "Projetar direção só com preço histórico",
     "Executado em 01/10/2026", d20["direcao"]["2. apenas preço histórico"],
     d20["majoritaria_teste"], None),
    (3, "Projetar direção com preço + notícia de jornal",
     "Reaproveitado (execução anterior, janela de teste diferente)",
     d18["direcao"]["Data Fusion: preços + notícia  [O QUE TÍNHAMOS ANTES]"],
     d18["majoritaria_teste"], None),
    (4, "Projetar direção com preço + TODOS os comunicados da CVM",
     "Executado em 01/10/2026", d20["direcao"]["4. preço + CVM (TODOS os comunicados)"],
     d20["majoritaria_teste"], None),
    (5, "Projetar direção com preço + SÓ Fato Relevante",
     "Executado em 01/10/2026", d20["direcao"]["5. preço + CVM (SÓ Fato Relevante)"],
     d20["majoritaria_teste"], None),
    (6, "Comparar SÓ Fato Relevante × TODOS os comunicados",
     "Executado em 01/10/2026 — ver aba McNemar (empate técnico)", None, None, None),
    (7, "Combinar notícia de jornal + SÓ Fato Relevante",
     "BLOQUEADO — precisa do corpus de notícias classificado, ausente nesta máquina",
     None, None, None),
]

r = 2
linha0 = r
for num, pedido, status, metricas, maj, _ in ITENS:
    texto_cel(ws1, r, 1, num, bold=True)
    texto_cel(ws1, r, 2, pedido)
    texto_cel(ws1, r, 3, status)
    if metricas:
        svm, xgb = metricas["SVM-RBF"], metricas["XGBoost"]
        melhor = svm if svm["acuracia"] >= xgb["acuracia"] else xgb
        input_cel(ws1, r, 4, f"{svm['acuracia']*100:.2f}% / {xgb['acuracia']*100:.2f}%", azul=False)
        ws1.cell(row=r, column=4).alignment = Alignment(horizontal="center")
        input_cel(ws1, r, 5, maj, "0.00%")
        input_cel(ws1, r, 6, melhor["p_binomial_vs_majoritaria"], "0.0000")
        ws1.cell(row=r, column=7, value=f"=IF(F{r}<0.05,\"Sim\",\"Não\")")
        ws1.cell(row=r, column=7).font = Font(name=FONTE_NOME, size=10)
        ws1.cell(row=r, column=7).alignment = Alignment(horizontal="center")
        if melhor["p_binomial_vs_majoritaria"] >= 0.05:
            for c in range(1, 8):
                pass  # sem destaque — nenhum item passa; destaque reservado se algum passar
    else:
        for c in range(4, 8):
            texto_cel(ws1, r, c, "—")
            ws1.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    for c in range(1, 8):
        ws1.cell(row=r, column=c).border = BORDA
        ws1.cell(row=r, column=c).alignment = Alignment(
            wrap_text=True, vertical="center",
            horizontal=ws1.cell(row=r, column=c).alignment.horizontal or "left")
    linha_zebra(ws1, r, linha0)
    ws1.row_dimensions[r].height = 60
    r += 1

r += 1
texto_cel(ws1, r, 1, "Em uma frase:", bold=True)
r += 1
ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
texto_cel(ws1, r, 1,
          ("Nenhum braço de direção (só preço, preço+CVM completo, preço+só Fato Relevante, "
           "preço+notícia) se distingue do palpite fixo. Item 6: usar só Fato Relevante ou "
           "usar todos os comunicados dá empate técnico (ver aba McNemar). Itens 1 e 7 não "
           "entram nesse veredito: o 1 é medição, não projeção; o 7 está bloqueado."))
ws1.cell(row=r, column=1).font = Font(name=FONTE_NOME, size=10, italic=True, color="595959")
ws1.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
ws1.row_dimensions[r].height = 40

# ═══ ABA 2 — DIREÇÃO: MÉTRICAS COMPLETAS ═══════════════════════════════════
ws2 = wb.create_sheet("Direção - métricas completas")
titulos2 = ["Braço", "Fonte", "Modelo", "N teste", "Classe\nmajoritária",
            "Acurácia", "Ganho\n(pontos)", "Precisão", "F1", "AUC",
            "p vs\nmajoritária", "Passa\n(p<0,05)?"]
larg2 = [38, 17, 10, 9, 12, 10, 10, 10, 9, 9, 11, 10]
cab(ws2, titulos2, larg2)

BRACOS = [
    ("Apenas preço histórico (item 2)", FONTE_20_CURTA, FONTE_20,
     d20["direcao"]["2. apenas preço histórico"], d20["majoritaria_teste"]),
    ("Preço + notícia de jornal (item 3)", FONTE_18_CURTA, FONTE_18,
     d18["direcao"]["Data Fusion: preços + notícia  [O QUE TÍNHAMOS ANTES]"], d18["majoritaria_teste"]),
    ("Preço + CVM, TODOS os comunicados (item 4)", FONTE_20_CURTA, FONTE_20,
     d20["direcao"]["4. preço + CVM (TODOS os comunicados)"], d20["majoritaria_teste"]),
    ("Preço + CVM, SÓ Fato Relevante (item 5)", FONTE_20_CURTA, FONTE_20,
     d20["direcao"]["5. preço + CVM (SÓ Fato Relevante)"], d20["majoritaria_teste"]),
]

r, linha0 = 2, 2
for nome, fonte_curta, fonte, metricas, maj in BRACOS:
    for mn in ["SVM-RBF", "XGBoost"]:
        m = metricas[mn]
        texto_cel(ws2, r, 1, nome)
        texto_cel(ws2, r, 2, fonte_curta)
        ws2.cell(row=r, column=2).comment = Comment(fonte, "gerador")
        texto_cel(ws2, r, 3, mn)
        input_cel(ws2, r, 4, m["n_teste"], "0")
        input_cel(ws2, r, 5, maj, "0.00%")
        input_cel(ws2, r, 6, m["acuracia"], "0.00%")
        ws2.cell(row=r, column=7, value=f"=(F{r}-E{r})*100")
        ws2.cell(row=r, column=7).number_format = "+0.00;-0.00"
        ws2.cell(row=r, column=7).font = Font(name=FONTE_NOME, size=10)
        input_cel(ws2, r, 8, m["precisao"], "0.00%")
        input_cel(ws2, r, 9, m["f1"], "0.00%")
        input_cel(ws2, r, 10, m["auc"], "0.000")
        input_cel(ws2, r, 11, m["p_binomial_vs_majoritaria"], "0.0000")
        ws2.cell(row=r, column=12, value=f"=IF(K{r}<0.05,\"Sim\",\"Não\")")
        ws2.cell(row=r, column=12).font = Font(name=FONTE_NOME, size=10)
        for c in range(4, 13):
            ws2.cell(row=r, column=c).alignment = Alignment(horizontal="right")
        for c in range(1, 13):
            ws2.cell(row=r, column=c).border = BORDA
        linha_zebra(ws2, r, linha0)
        r += 1
ws2.auto_filter.ref = f"A1:{get_column_letter(12)}{r-1}"

r += 1
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
texto_cel(ws2, r, 1,
          ("Atenção: 'Preço + notícia' usa uma janela de teste diferente dos outros três "
           "braços (653 pregões, execução anterior, x 579 pregões, execução de hoje) — "
           "compare só de forma descritiva, não como se fosse o mesmo teste pareado."))
ws2.cell(row=r, column=1).font = Font(name=FONTE_NOME, size=9, italic=True, color="595959")
ws2.cell(row=r, column=1).alignment = Alignment(wrap_text=True)

# ═══ ABA 3 — McNEMAR (pares) ════════════════════════════════════════════════
ws3 = wb.create_sheet("McNemar - pares")
titulos3 = ["Comparação", "Modelo", "Ganhou\n(só o novo acertou)",
            "Perdeu\n(só o antigo acertou)", "p-valor", "Significativo\n(p<0,05)?"]
larg3 = [56, 10, 13, 15, 10, 12]
cab(ws3, titulos3, larg3)

NOMES_PAR = {
    "5 vs 2": "5. Preço+FR acrescenta sobre 2. só preço?",
    "4 vs 2": "4. Preço+todos os comunicados acrescenta sobre 2. só preço?",
    "6": "6. SÓ Fato Relevante × TODOS os comunicados (item 6)",
}
r, linha0 = 2, 2
for chave, m in d20["mcnemar"].items():
    modelo = chave.split(" ")[0]
    # identifica o par pelo texto completo salvo no JSON
    if "5 vs 2" in chave:
        rotulo = NOMES_PAR["5 vs 2"]
    elif "4 vs 2" in chave:
        rotulo = NOMES_PAR["4 vs 2"]
    else:
        rotulo = NOMES_PAR["6"]
    texto_cel(ws3, r, 1, rotulo)
    texto_cel(ws3, r, 2, modelo)
    input_cel(ws3, r, 3, m["acertos_so_do_novo"], "0", azul=False)
    input_cel(ws3, r, 4, m["acertos_so_do_antigo"], "0", azul=False)
    input_cel(ws3, r, 5, m["p"], "0.0000")
    ws3.cell(row=r, column=6, value=f"=IF(E{r}<0.05,\"Sim\",\"Não\")")
    ws3.cell(row=r, column=6).font = Font(name=FONTE_NOME, size=10)
    for c in range(3, 7):
        ws3.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    for c in range(1, 7):
        ws3.cell(row=r, column=c).border = BORDA
    linha_zebra(ws3, r, linha0)
    r += 1
ws3.auto_filter.ref = f"A1:F{r-1}"

r += 1
ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
texto_cel(ws3, r, 1,
          "Fonte: CVM/dados/preco_e_cvm_sem_noticia.json — executado em 01/10/2026, PETR4.")
ws3.cell(row=r, column=1).font = Font(name=FONTE_NOME, size=9, italic=True, color="595959")

# ═══ ABA 4 — NOTAS E LIMITAÇÕES ═════════════════════════════════════════════
ws4 = wb.create_sheet("Notas e limitações")
ws4.column_dimensions["A"].width = 110
NOTAS = [
    ("Protocolo", "Capítulo 3 da dissertação: SVM-RBF e XGBoost (300 árvores, "
     "profundidade 3, taxa 0,05, subamostragem 0,9), divisão cronológica 60/15/25 "
     "(treino/validação/teste), teste binomial contra a classe majoritária e "
     "teste de McNemar entre pares de modelo."),
    ("Itens 2, 4, 5 e 6", "Executados em 01/10/2026 nesta máquina "
     "(CVM/20_preco_e_cvm_sem_noticia.py). O GARCH(1,1) foi reajustado aqui "
     "(ω=0,145 α=0,078 β=0,899 — bate com o α=0,085/β=0,899 já registrado antes), "
     "porque Mestrado_PETR4/base_master_petr4.csv não está nesta máquina."),
    ("Item 3", "NÃO foi reexecutado hoje. O número vem de "
     "CVM/dados/pesquisa_com_cvm.json, calculado antes, numa janela de teste "
     "diferente (653 pregões, terminando mais cedo, contra 579 pregões de hoje, "
     "terminando em 15/09/2026). O p-valor contra a classe majoritária foi "
     "recalculado agora (o arquivo original só trazia p contra 50%)."),
    ("Item 7", "BLOQUEADO. Pede notícia de jornal combinada só com Fato Relevante. "
     "O script já existe (CVM/19_fr_vs_todos_comunicados.py) mas depende de "
     "Mestrado_PETR4/base_master_petr4.csv, que por sua vez depende do corpus de "
     "notícias classificado (indice_sentimento_petr4.csv) — nenhum dos dois está "
     "nesta máquina. Para fechar, traga esses dois arquivos de onde estão, ou "
     "rode o script 19 lá."),
    ("Cor azul", "Toda célula em azul é um resultado medido (saída de um modelo "
     "treinado ou de um teste estatístico) — não é fórmula, é dado de origem, "
     "com a fonte anotada ao lado ou em comentário."),
    ("Cor preta / fórmula", "'Ganho (pontos)' e 'Passa (p<0,05)?' são sempre "
     "fórmula, recalculam se você trocar a acurácia ou o p-valor ao lado."),
    ("Arredondamento em cadeia", "'Ganho (pontos)' é calculado aqui a partir da "
     "acurácia e da classe majoritária já arredondadas em 4 casas no arquivo de "
     "origem — por isso pode diferir em até 0,01 ponto percentual do número "
     "'ganho_pp_vs_majoritaria' registrado originalmente pelo script (ex.: "
     "aqui dá -1,39; o script registrou -1,38). A diferença não muda nenhuma "
     "conclusão; é só precisão de exibição."),
]
r = 1
for titulo, texto in NOTAS:
    texto_cel(ws4, r, 1, titulo, bold=True)
    ws4.cell(row=r, column=1).font = Font(name=FONTE_NOME, size=11, bold=True, color=AZUL_CAB)
    r += 1
    texto_cel(ws4, r, 1, texto)
    ws4.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.row_dimensions[r].height = 48
    r += 2

wb.save(SAIDA)
print(f"gravado: {SAIDA}")
